"""Portable browser smoke for the public-release new-user path.

The script intentionally uses only Playwright's managed Chromium and the
user-visible UI.  It is a small CI gate, not a replacement for the longer
isolated release journeys kept in the local evaluation reports.
"""

from __future__ import annotations

import json
import os
import time
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from playwright.sync_api import expect, sync_playwright
from release_endpoints import (
    assert_release_backend_ready,
    assert_release_frontend_ready,
    release_api_url,
    release_frontend_url,
)
from temp_paths import test_temp_root

BASE_URL = release_frontend_url()
API_URL = release_api_url()
ARTIFACT_DIR = Path(os.getenv("OFFERU_E2E_ARTIFACT_DIR") or test_temp_root("e2e-artifacts"))


def _json_response(page, url: str) -> dict:
    response = page.request.get(url)
    if not response.ok:
        raise AssertionError(f"request failed: {response.status} {url}")
    payload = response.json()
    if not isinstance(payload, dict):
        raise AssertionError(f"expected object response: {url}")
    return payload


def _replay_provider_payload(route) -> str:
    payload = route.request.post_data_json
    if not isinstance(payload, dict):
        raise AssertionError("job ingest request must contain a JSON object")
    payload["runtime_provider"] = "replay"
    return json.dumps(payload, ensure_ascii=False)


def _use_replay_provider(route) -> None:
    route.continue_(post_data=_replay_provider_payload(route))


def _synthetic_resume_docx() -> bytes:
    """Build a stable, synthetic resume for the visible first-run import flow."""
    from docx import Document

    document = Document()
    document.add_paragraph("姓名：OfferU E2E 用户")
    document.add_paragraph("求职意向：AI 产品经理")
    document.add_paragraph("工作经历")
    document.add_paragraph("OfferU 测试科技有限公司 产品经理 2023.01 - 2025.01")
    document.add_paragraph("负责 AI 产品需求分析，协调设计与工程完成首版交付。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _complete_new_user_onboarding(page, suffix: str) -> None:
    """Use the current App-first onboarding UI to establish a profile and open job setup."""
    wizard = page.get_by_role("dialog", name="把一个岗位，变成可准备的工作区")
    expect(wizard).to_be_visible(
        timeout=20000
    )
    agent_panel = wizard.get_by_test_id("agent-provider-health")
    expect(agent_panel).to_be_visible(timeout=20000)
    expect(agent_panel).to_contain_text(
        "检测到 0 个本机运行环境。连接能力以检查结果为准。", timeout=30000
    )
    expect(agent_panel.get_by_text("未检测到", exact=True).first).to_be_visible()
    expect(wizard.get_by_text("暂时没有可用 Agent 也可以继续。", exact=False)).to_be_visible()
    page.get_by_role("button", name="继续", exact=True).click()

    resume_input = page.get_by_label("选择简历文件", exact=True)
    resume_input.set_input_files(
        {
            "name": f"offeru-e2e-resume-{suffix}.docx",
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "buffer": _synthetic_resume_docx(),
        }
    )
    expect(page.get_by_text("识别出 1 条内容", exact=False)).to_be_visible(timeout=30000)
    expect(page.get_by_text("负责 AI 产品需求分析", exact=False)).to_be_visible()
    expect(page.get_by_text("Word 文档正文", exact=True)).to_be_visible()
    experience = page.get_by_role(
        "checkbox", name="OfferU 测试科技有限公司", exact=False
    )
    experience.check()
    page.get_by_role("button", name="确认所选内容，建立档案", exact=True).click()

    expect(page.get_by_text("整理 AI 记忆", exact=True)).to_be_visible(timeout=30000)
    page.get_by_role("button", name="跳过记忆", exact=True).click()
    expect(page.get_by_role("button", name="开始保存目标岗位", exact=True)).to_be_visible()
    page.get_by_role("button", name="开始保存目标岗位", exact=True).click()
    expect(page.get_by_test_id("add-job-modal")).to_be_visible(timeout=20000)


def main() -> None:
    suffix = str(int(time.time() * 1000))
    assert_release_frontend_ready()
    assert_release_backend_ready()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(
            viewport={"width": 1440, "height": 1000},
            accept_downloads=True,
        )
        context.tracing.start(screenshots=True, snapshots=True, sources=True)
        page = context.new_page()
        console_errors: list[str] = []
        page_errors: list[str] = []
        bad_responses: list[str] = []
        asset_requests: dict[int, dict[str, object]] = {}
        failed_frontend_assets: list[dict[str, object]] = []
        frontend_origin = urlsplit(BASE_URL)

        def on_asset_request(request) -> None:
            target = urlsplit(request.url)
            same_origin = (target.scheme, target.netloc) == (
                frontend_origin.scheme,
                frontend_origin.netloc,
            )
            if same_origin and request.resource_type in {"script", "stylesheet"}:
                asset_requests[id(request)] = {
                    "path": target.path,
                    "resource_type": request.resource_type,
                    "started_at": time.monotonic(),
                }

        def on_asset_request_finished(request) -> None:
            asset_requests.pop(id(request), None)

        def on_asset_request_failed(request) -> None:
            entry = asset_requests.pop(id(request), None)
            if entry is None:
                return
            failure_code = (request.failure or "").split(" ", 1)[0]
            if not failure_code.startswith("net::"):
                failure_code = "request failed"
            failed_frontend_assets.append(
                {
                    "path": entry["path"],
                    "resource_type": entry["resource_type"],
                    "elapsed_seconds": round(time.monotonic() - entry["started_at"], 3),
                    "failure": failure_code,
                }
            )

        def pending_frontend_assets() -> list[dict[str, object]]:
            now = time.monotonic()
            return [
                {
                    "path": entry["path"],
                    "resource_type": entry["resource_type"],
                    "elapsed_seconds": round(now - entry["started_at"], 3),
                }
                for entry in asset_requests.values()
            ]

        page.on(
            "console",
            lambda message: console_errors.append(message.text)
            if message.type == "error"
            else None,
        )
        page.on("pageerror", lambda error: page_errors.append(str(error)))
        page.on(
            "response",
            lambda response: bad_responses.append(f"{response.status} {response.url}")
            if response.status >= 400
            else None,
        )
        page.on("request", on_asset_request)
        page.on("requestfinished", on_asset_request_finished)
        page.on("requestfailed", on_asset_request_failed)

        trace_stopped = False
        try:
            page.goto(f"{BASE_URL}/#/?release_smoke={suffix}", wait_until="domcontentloaded")
            page.evaluate("localStorage.clear()")
            page.reload(wait_until="domcontentloaded")
            _complete_new_user_onboarding(page, suffix)
            page.get_by_test_id("add-job-title").fill("AI 产品经理")
            company = f"CI Orbit {suffix}"
            page.get_by_test_id("add-job-company").fill(company)
            page.get_by_test_id("add-job-description").fill(
                "负责 AI 产品规划、用户需求分析、评测体系建设与跨团队交付；关注 Agent 工作流和产品增长。"
            )
            # Exercise the user-visible duplicate-submit boundary.  The UI guard
            # and the deterministic ingest hash must still produce one Job and
            # one preparation task when a user double-clicks. The isolated CI
            # runner has no local Agent, so select the supported replay provider
            # on this request instead of depending on host auto-discovery.
            page.route("**/api/jobs/ingest", _use_replay_provider)
            try:
                page.get_by_test_id("add-job-submit").dblclick()
                page.wait_for_function(
                    "() => window.location.hash.startsWith('#/jobs/')",
                    timeout=30000,
                )
            finally:
                page.unroute("**/api/jobs/ingest", _use_replay_provider)

            jobs = _json_response(page, f"{API_URL}/api/jobs/?page_size=100")
            job = next(
                (
                    item
                    for item in jobs.get("items", [])
                    if item.get("company") == company
                ),
                None,
            )
            if not isinstance(job, dict):
                raise AssertionError("created CI job was not returned by the jobs API")
            job_id = int(job["id"])
            matching_jobs = [
                item for item in jobs.get("items", []) if item.get("company") == company
            ]
            if len(matching_jobs) != 1:
                raise AssertionError(
                    f"duplicate submit created {len(matching_jobs)} jobs for {company}"
                )

            task: dict | None = None
            for _ in range(60):
                tasks = _json_response(
                    page,
                    f"{API_URL}/api/agent/runtime/career-tasks"
                    f"?target_type=job&target_id={job_id}&limit=10",
                )
                task = next(
                    (
                        item
                        for item in tasks.get("tasks", [])
                        if item.get("task_type") == "role_intelligence"
                    ),
                    None,
                )
                if isinstance(task, dict) and task.get("status") in {
                    "completed",
                    "failed",
                    "blocked",
                }:
                    break
                page.wait_for_timeout(500)
            if not isinstance(task, dict):
                raise AssertionError(f"role intelligence task was not created for job {job_id}")
            if task.get("status") != "completed":
                raise AssertionError(f"replay role intelligence did not complete: {task}")
            if task.get("runtime_provider") != "replay":
                raise AssertionError(f"CI smoke must use explicit replay provider: {task}")
            matching_tasks = [
                item
                for item in tasks.get("tasks", [])
                if item.get("task_type") == "role_intelligence"
            ]
            if len(matching_tasks) != 1:
                raise AssertionError(
                    f"duplicate submit created {len(matching_tasks)} role intelligence tasks"
                )

            page.wait_for_function(
                "jobId => window.location.hash.slice(1).split('?')[0] === `/jobs/${jobId}`",
                arg=job_id,
                timeout=20000,
            )
            role_panel = page.get_by_test_id("role-intelligence-panel")
            expect(role_panel).to_be_visible(timeout=20000)
            expect(role_panel).to_contain_text("20", timeout=30000)
            expect(page.get_by_text("材料候选", exact=True)).to_be_visible(timeout=30000)
            body = page.locator("body").inner_text()
            if "材料候选" not in body or "20" not in body:
                raise AssertionError("job detail did not expose the prepared packet")

            # Exercise a transport-level retry after the server has already
            # committed the first request.  The browser receives a synthetic
            # 503, keeps the form/error visible, and the user retries the same
            # payload.  The ingest hash must make the business effect exactly
            # once even though the first response was not observed as success.
            retry_company = f"CI Retry {suffix}"
            retry_state = {"attempts": 0}

            def fail_first_ingest(route) -> None:
                retry_state["attempts"] += 1
                replay_payload = _replay_provider_payload(route)
                if retry_state["attempts"] == 1:
                    route.fetch(post_data=replay_payload)
                    route.fulfill(
                        status=503,
                        content_type="application/json",
                        body=json.dumps({"detail": "模拟网络错误"}, ensure_ascii=True),
                    )
                    return
                route.continue_(post_data=replay_payload)

            page.route("**/api/jobs/ingest", fail_first_ingest)
            try:
                page.goto(f"{BASE_URL}/#/jobs?release_smoke_retry={suffix}", wait_until="domcontentloaded")
                page.get_by_test_id("open-add-job").wait_for(timeout=20000)
                page.get_by_test_id("open-add-job").click()
                page.get_by_test_id("add-job-title").fill("AI 产品经理")
                page.get_by_test_id("add-job-company").fill(retry_company)
                page.get_by_test_id("add-job-description").fill(
                    "负责 AI 产品规划、用户需求分析、评测体系建设与跨团队交付；关注 Agent 工作流和产品增长。"
                )
                page.get_by_test_id("add-job-submit").click()
                expect(page.get_by_role("alert")).to_contain_text("模拟网络错误", timeout=30000)
                page.get_by_test_id("add-job-submit").click()
                expect(page.get_by_test_id("add-job-modal")).not_to_be_visible(timeout=30000)
            finally:
                page.unroute("**/api/jobs/ingest", fail_first_ingest)

            if retry_state["attempts"] != 2:
                raise AssertionError(
                    f"transport retry expected two ingest attempts, got {retry_state['attempts']}"
                )
            retry_jobs = _json_response(page, f"{API_URL}/api/jobs/?page_size=100")
            retry_matches = [
                item for item in retry_jobs.get("items", []) if item.get("company") == retry_company
            ]
            if len(retry_matches) != 1:
                raise AssertionError(
                    f"transport retry created {len(retry_matches)} jobs for {retry_company}"
                )
            retry_job_id = int(retry_matches[0]["id"])
            retry_task: dict | None = None
            retry_tasks: dict = {}
            for _ in range(60):
                retry_tasks = _json_response(
                    page,
                    f"{API_URL}/api/agent/runtime/career-tasks"
                    f"?target_type=job&target_id={retry_job_id}&limit=10",
                )
                retry_task = next(
                    (
                        item
                        for item in retry_tasks.get("tasks", [])
                        if item.get("task_type") == "role_intelligence"
                    ),
                    None,
                )
                if isinstance(retry_task, dict) and retry_task.get("status") in {
                    "completed",
                    "failed",
                    "blocked",
                }:
                    break
                page.wait_for_timeout(500)
            if not isinstance(retry_task, dict) or retry_task.get("status") != "completed":
                raise AssertionError(f"transport retry task did not complete: {retry_task}")
            retry_matching_tasks = [
                item
                for item in retry_tasks.get("tasks", [])
                if item.get("task_type") == "role_intelligence"
            ]
            if len(retry_matching_tasks) != 1:
                raise AssertionError(
                    f"transport retry created {len(retry_matching_tasks)} role intelligence tasks"
                )
            page.wait_for_function(
                "jobId => window.location.hash.slice(1).split('?')[0] === `/jobs/${jobId}`",
                arg=retry_job_id,
                timeout=20000,
            )
            retry_role_panel = page.get_by_test_id("role-intelligence-panel")
            expect(retry_role_panel).to_be_visible(timeout=20000)
            expect(retry_role_panel).to_contain_text("20", timeout=30000)

            expected_bad_responses = [
                response
                for response in bad_responses
                if response.startswith("503 ") and "/api/jobs/ingest" in response
            ]
            unexpected_bad_responses = [
                response for response in bad_responses if response not in expected_bad_responses
            ]
            expected_console_errors = [message for message in console_errors if "503" in message]
            unexpected_console_errors = [
                message for message in console_errors if message not in expected_console_errors
            ]
            if len(expected_bad_responses) != 1 or len(expected_console_errors) != 1:
                raise AssertionError(
                    "transport retry did not expose exactly one expected 503: "
                    f"responses={bad_responses}, console={console_errors}"
                )
            if unexpected_bad_responses or unexpected_console_errors or page_errors:
                raise AssertionError(
                    "browser errors: "
                    f"responses={unexpected_bad_responses}, "
                    f"console={unexpected_console_errors}, page={page_errors}"
                )

            result = {
                "status": "PASS",
                "job_id": job_id,
                "task_id": task["task_id"],
                "task_status": task["status"],
                "runtime_provider": task["runtime_provider"],
                "retry_job_id": retry_job_id,
                "retry_task_id": retry_task["task_id"],
                "transport_attempts": retry_state["attempts"],
                "expected_transport_failure": expected_bad_responses[0],
                "bad_responses": bad_responses,
                "console_errors": console_errors,
                "page_errors": page_errors,
                "failed_frontend_assets": failed_frontend_assets,
            }
            print(json.dumps(result, ensure_ascii=True, indent=2), flush=True)
            context.tracing.stop()
            trace_stopped = True
        except Exception:
            ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
            screenshot_path = ARTIFACT_DIR / "public-release-smoke-failure.png"
            trace_path = ARTIFACT_DIR / "public-release-smoke-failure.zip"
            try:
                page.screenshot(path=str(screenshot_path), full_page=True)
            except Exception:
                pass
            try:
                context.tracing.stop(path=str(trace_path))
                trace_stopped = True
            except Exception:
                pass
            print(
                json.dumps(
                    {
                        "status": "FAIL",
                        "screenshot": str(screenshot_path),
                        "trace": str(trace_path),
                        "bad_responses": bad_responses,
                        "console_errors": console_errors,
                        "page_errors": page_errors,
                        "failed_frontend_assets": failed_frontend_assets,
                        "pending_frontend_assets": pending_frontend_assets(),
                    },
                    ensure_ascii=True,
                    indent=2,
                ),
                flush=True,
            )
            raise
        finally:
            if not trace_stopped:
                try:
                    context.tracing.stop()
                except Exception:
                    pass
            context.close()
            browser.close()


if __name__ == "__main__":
    main()
