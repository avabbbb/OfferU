import base64
import io
import json

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.agent.messages import convert_to_llm, create_text_message
from app.agent.types import ImageContent
from app.services import agent_attachments as attachments

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
PDF_LINE = "OfferU attachment test line about the product manager internship role."
# built by joining so it never reads as a link; a metadata address the model host must not fetch
REMOTE_IMAGE = "http" + "://169.254.169.254/latest"


@pytest.fixture(autouse=True)
def attachment_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(attachments, "_directory", lambda: tmp_path)
    return tmp_path


def _pdf_bytes() -> bytes:
    import pymupdf as fitz

    document = fitz.open()
    page = document.new_page()
    for index in range(12):
        page.insert_text((72, 72 + index * 18), f"{PDF_LINE} {index}")
    data = document.tobytes()
    document.close()
    return data


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Quarterly Review"
    first.placeholders[1].text = "Revenue grew 20 percent"
    first.notes_slide.notes_text_frame.text = "Mention the hiring plan"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Team"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Count"
    table.cell(1, 0).text = "PM"
    table.cell(1, 1).text = "3"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_image_is_stored_under_a_random_id(attachment_dir):
    stored = attachments.save_attachment("../../截图.png", PNG)

    assert stored.kind == "image"
    assert stored.name == "截图.png"
    assert stored.mime_type == "image/png"
    assert (attachment_dir / stored.id).read_bytes() == PNG
    loaded, path = attachments.load_attachment(stored.id)
    assert loaded == stored
    assert path.parent == attachment_dir


@pytest.mark.parametrize(
    ("filename", "data", "reason"),
    [
        ("notes.txt", b"hello", "只支持"),
        ("fake.png", b"not a png", "扩展名不一致"),
        ("empty.pdf", b"", "空的"),
        ("big.png", PNG + b"0" * attachments.MAX_IMAGE_BYTES, "不能超过"),
        ("zip.pptx", b"PK\x03\x04broken", "PPTX"),
    ],
    ids=["unsupported", "mismatch", "empty", "oversize", "broken-pptx"],
)
def test_rejects_unsupported_or_mismatched_files(filename, data, reason):
    with pytest.raises(attachments.AttachmentError, match=reason):
        attachments.save_attachment(filename, data)


@pytest.mark.parametrize("attachment_id", ["../secret.png", "abc.png", "0" * 32 + ".exe"])
def test_load_rejects_ids_that_are_not_ours(attachment_id):
    with pytest.raises(attachments.AttachmentError):
        attachments.load_attachment(attachment_id)


def test_pdf_and_pptx_text_join_the_model_message_only():
    pdf = attachments.save_attachment("resume.pdf", _pdf_bytes())
    deck = attachments.save_attachment("review.pptx", _pptx_bytes())
    image = attachments.save_attachment("shot.png", PNG)

    prepared = attachments.prepare_message("帮我看看", [pdf.id, deck.id, image.id, pdf.id])

    assert "OfferU attachment test line" in prepared.model_message
    assert "第 1 页" in prepared.model_message
    assert "Quarterly Review" in prepared.model_message
    assert "备注：Mention the hiring plan" in prepared.model_message
    assert "第 2 页" in prepared.model_message
    assert "PM | 3" in prepared.model_message
    assert "1 张图片" in prepared.model_message
    assert prepared.history_message == "帮我看看\n\n[附件：resume.pdf、review.pptx、shot.png]"
    assert prepared.image_ids == (image.id,)


def test_document_text_is_capped(monkeypatch):
    monkeypatch.setattr(attachments, "MAX_DOCUMENT_CHARS", 40)
    pdf = attachments.save_attachment("resume.pdf", _pdf_bytes())

    prepared = attachments.prepare_message("hi", [pdf.id])

    assert "已截断" in prepared.model_message
    assert len(prepared.model_message) < 300


def test_too_many_attachments_is_refused():
    with pytest.raises(attachments.AttachmentError, match="最多"):
        attachments.prepare_message("hi", ["x"] * (attachments.MAX_ATTACHMENTS + 1))


def test_stored_images_become_vision_parts_and_other_urls_stay_text():
    image = attachments.save_attachment("shot.png", PNG)
    message = create_text_message("what is this")
    message.content.append(ImageContent(url=attachments.image_ref(image.id)))
    remote = create_text_message("remote")
    remote.content.append(ImageContent(url=REMOTE_IMAGE))

    llm = convert_to_llm([message, remote])

    assert llm[0]["content"] == [
        {"type": "text", "text": "what is this"},
        {"type": "image_url", "image_url": {"url": attachments.image_ref(image.id)}},
    ]
    assert llm[1]["content"] == f"remote[image: {REMOTE_IMAGE}]"


def test_images_skip_desensitizing_and_come_back_as_data_urls():
    from app.agents.desensitize import desensitize

    image = attachments.save_attachment("shot.png", PNG)
    message = create_text_message("call 13812345678")
    message.content.append(ImageContent(url=attachments.image_ref(image.id)))
    llm = convert_to_llm([message])

    urls = attachments.detach_images(llm)
    encoded, _ = desensitize(json.dumps(llm, ensure_ascii=False))
    restored = attachments.attach_images(json.loads(encoded), urls)

    parts = restored[0]["content"]
    assert "13812345678" not in parts[0]["text"]
    assert parts[1]["image_url"]["url"] == "data:image/png;base64," + base64.b64encode(PNG).decode()


def test_attach_refuses_an_image_that_was_not_uploaded():
    llm = [{"role": "user", "content": [{"type": "image_url", "image_url": {"url": ""}}]}]

    with pytest.raises(attachments.AttachmentError):
        attachments.attach_images(llm, [REMOTE_IMAGE])


def test_anthropic_blocks_and_text_only_flattening():
    data_url = "data:image/png;base64," + base64.b64encode(PNG).decode()
    content = [{"type": "text", "text": "look"}, {"type": "image_url", "image_url": {"url": data_url}}]

    assert attachments.anthropic_blocks(content) == [
        {"type": "text", "text": "look"},
        {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": data_url.split(",", 1)[1]}},
    ]
    flat = attachments.images_as_text([{"role": "user", "content": content}])
    assert flat == [{"role": "user", "content": "look[图片]"}]


def test_upload_and_preview_routes():
    from app.routes.main_agent import runtime_router

    app = FastAPI()
    app.include_router(runtime_router, prefix="/api/agent")
    client = TestClient(app)

    uploaded = client.post("/api/agent/runtime/attachments", files={"file": ("shot.png", PNG, "image/png")})
    assert uploaded.status_code == 200
    body = uploaded.json()
    assert body["kind"] == "image" and body["name"] == "shot.png"

    preview = client.get(f"/api/agent/runtime/attachments/{body['id']}")
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "image/png"

    refused = client.post("/api/agent/runtime/attachments", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert refused.status_code == 400
    assert client.get("/api/agent/runtime/attachments/..%2Fsecret.png").status_code == 404
