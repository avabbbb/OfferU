"""主线：Agent（MCP）和工作台（HTTP）用同一套工具写同一份数据，写了直接生效、能撤销，今天页给出下一步。"""

import asyncio
import json
from datetime import date

from app.core.tools import call, overview


def test_core_main_path(tmp_path, monkeypatch):
    monkeypatch.setenv("OFFERU_DATA_DIR", str(tmp_path))
    agent = "codex"
    assert overview(today=date(2026, 10, 20))["next_step"]["title"] == "先建档案"
    call("save", {"type": "profile", "data": {"name": "张三"}}, actor="你")
    assert overview(today=date(2026, 10, 20))["next_step"]["title"] == "写一份简历"
    call("save", {"type": "doc", "data": {"kind": "resume", "title": "主简历", "body": "# 张三"}}, actor="你")
    assert overview(today=date(2026, 10, 20))["next_step"]["title"] == "存几个想投的岗位"

    job = call("save", {"type": "job", "data": {"company": "百度", "title": "AI 产品经理"}}, actor=agent)["record"]
    assert job["status"] == "saved"
    assert overview(today=date(2026, 10, 20))["next_step"]["title"] == "投递「百度 · AI 产品经理」"

    call("save", {"type": "event", "data": {"title": "百度笔试", "kind": "test", "at": "2026-10-21T19:00", "job_id": job["id"]}}, actor=agent)
    view = overview(today=date(2026, 10, 20))
    assert view["next_step"]["title"] == "准备「百度笔试」" and view["todos"]["week"][0]["job_label"] == "百度 · AI 产品经理"

    applied = call("save", {"type": "job", "id": job["id"], "data": {"status": "applied"}}, actor="你")
    assert applied["record"]["applied_at"] and overview()["stage"] == "applying"
    call("undo", {"change_id": applied["change_id"]}, actor="你")
    assert call("get", {"type": "job", "id": job["id"]}, actor="你")["record"]["status"] == "saved"

    gone = call("delete", {"type": "job", "id": job["id"]}, actor=agent)
    call("undo", {"change_id": gone["change_id"]}, actor="你")
    assert call("get", {"type": "job", "id": job["id"]}, actor="你")["events"][0]["title"] == "百度笔试"

    call("save", {"type": "memory", "data": {"kind": "experience", "title": "学生会外联部长", "body": "拉到 3 家企业赞助"}}, actor=agent)
    assert call("search", {"query": "赞助"}, actor=agent)["items"][0]["label"] == "学生会外联部长"
    assert {c["actor"] for c in call("changes", {}, actor="你")["items"]} == {"codex", "你"}


def test_mcp_and_http_share_tools(tmp_path, monkeypatch):
    monkeypatch.setenv("OFFERU_DATA_DIR", str(tmp_path))
    from fastapi.testclient import TestClient

    from app.main import app
    from app.mcp_server import mcp

    names = [t.name for t in asyncio.run(mcp.list_tools())]
    assert names[0] == "overview" and len(names) <= 15
    asyncio.run(mcp.call_tool("save", {"type": "job", "data": {"company": "米哈游"}}))
    with TestClient(app, base_url="http://127.0.0.1:8766") as client:
        listed = client.post("/api/tools/list", json={"type": "job"}).json()
        assert listed["items"][0]["company"] == "米哈游"
        bad = client.post("/api/tools/save", json={"type": "job", "data": {"status": "maybe"}})
        assert bad.status_code == 400 and "status" in json.dumps(bad.json(), ensure_ascii=False)


def test_resume_import_edit_export(tmp_path, monkeypatch):
    """导入旧简历 → 档案 → Markdown 简历（连续编辑只算一条改动）→ 编辑器和 PDF 用同一份模板样式。"""
    monkeypatch.setenv("OFFERU_DATA_DIR", str(tmp_path))
    import docx
    import fitz
    from fastapi.testclient import TestClient

    from app.main import app

    old = docx.Document()
    old.add_paragraph("李四 复旦大学 计算机 2025 届")
    old.add_paragraph("字节跳动 后端实习：接口延迟从 120ms 降到 40ms")
    old.save(tmp_path / "old.docx")
    with TestClient(app, base_url="http://127.0.0.1:8766") as client:
        name = client.post("/api/files", files={"file": ("我的简历.docx", (tmp_path / "old.docx").read_bytes())}).json()["name"]
        text = client.post("/api/tools/read_file", json={"path": name}).json()["text"]
        assert "120ms" in text

        body = "# 李四\n138 0000 0000 · li@example.com\n\n## 实习经历\n### 字节跳动 | 后端开发实习生 | 2024.06 – 2024.09\n- 接口延迟从 120ms 降到 40ms\n"
        doc = client.post("/api/tools/save", json={"type": "doc", "data": {"kind": "resume", "title": "主简历", "body": "# 李四", "template": "compact"}}).json()["record"]
        first = client.post("/api/tools/save", json={"type": "doc", "id": doc["id"], "data": {"body": "# 李四\n"}}).json()["change_id"]
        second = client.post("/api/tools/save", json={"type": "doc", "id": doc["id"], "data": {"body": body}}).json()["change_id"]
        assert first == second
        assert client.post("/api/tools/save", json={"type": "doc", "id": doc["id"], "data": {"template": "fancy"}}).status_code == 400

        assert "--body-size: 9.2pt" in client.get("/api/resume/templates/compact.css").text
        pdf = client.get(f"/api/resume/{doc['id']}/pdf").content
        assert "接口延迟从 120ms 降到 40ms" in fitz.open(stream=pdf, filetype="pdf")[0].get_text()
        path = client.post("/api/tools/export_resume_pdf", json={"id": doc["id"]}).json()["path"]
        assert path.endswith("主简历.pdf")

        client.post("/api/tools/undo", json={"change_id": second})
        assert client.post("/api/tools/get", json={"type": "doc", "id": doc["id"]}).json()["record"]["body"] == "# 李四"
