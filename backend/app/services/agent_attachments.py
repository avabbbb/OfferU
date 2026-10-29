"""Files the user attaches in the Agent composer.

Images reach the model as vision input; PDF and PPTX are read to text once, at
upload, and joined into the user's message when it is sent. Every file is stored
under a random id so a client filename never becomes part of a path, and the
image bytes are only inlined after the message has been desensitized.
"""

from __future__ import annotations

import base64
import io
import json
import re
import uuid
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

IMAGE_TYPES: dict[str, str] = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
DOCUMENT_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_UPLOAD_BYTES = max(MAX_IMAGE_BYTES, MAX_DOCUMENT_BYTES)
MAX_ATTACHMENTS = 8
MAX_PDF_PAGES = 60
MAX_SLIDES = 200
MAX_PPTX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCUMENT_CHARS = 30_000
MAX_TOTAL_DOCUMENT_CHARS = 60_000
IMAGE_REF_PREFIX = "offeru-attachment://"

_ID_RE = re.compile(r"^[0-9a-f]{32}\.(png|jpg|jpeg|webp|gif|pdf|pptx)$")
_CONTROL_RE = re.compile(r"[\x00-\x1f\x7f]")
_TRUNCATED = "\n…（内容过长，已截断）"


class AttachmentError(ValueError):
    """A user-facing reason an attachment cannot be accepted or used."""


@dataclass(frozen=True)
class StoredAttachment:
    id: str
    name: str
    kind: str  # image | document
    mime_type: str
    size: int

    def public(self) -> dict[str, Any]:
        return asdict(self)


def _directory() -> Path:
    from app.core import store

    return store.data_dir("agent_attachments")


def _display_name(filename: str) -> str:
    name = _CONTROL_RE.sub("", Path(filename.replace("\\", "/")).name).strip()
    return (name or "未命名文件")[:120]


def _matches_signature(suffix: str, data: bytes) -> bool:
    if suffix == ".png":
        return data.startswith(b"\x89PNG\r\n\x1a\n")
    if suffix in {".jpg", ".jpeg"}:
        return data.startswith(b"\xff\xd8\xff")
    if suffix == ".gif":
        return data.startswith((b"GIF87a", b"GIF89a"))
    if suffix == ".webp":
        return data[:4] == b"RIFF" and data[8:12] == b"WEBP"
    if suffix == ".pdf":
        return data[:1024].lstrip().startswith(b"%PDF-")
    if suffix == ".pptx":
        return data.startswith(b"PK\x03\x04")
    return False


def _pdf_text(data: bytes) -> str:
    import pymupdf as fitz

    from app.services.resume_parser import _parse_pdf_document_sync

    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise AttachmentError("这个 PDF 有密码保护，请先取消密码再上传")
            if document.page_count > MAX_PDF_PAGES:
                raise AttachmentError(f"PDF 最多支持 {MAX_PDF_PAGES} 页，这个文件有 {document.page_count} 页")
    except AttachmentError:
        raise
    except Exception as exc:
        raise AttachmentError("这个 PDF 无法打开，可能已损坏") from exc
    try:
        return _parse_pdf_document_sync(data).text.strip()
    except Exception as exc:
        raise AttachmentError("这个 PDF 的文字读取失败") from exc


def _shape_lines(shape: Any) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        lines: list[str] = []
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        rows: list[str] = []
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return rows
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        return [text] if text else []
    return []


def _pptx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            total = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise AttachmentError("这个 PPTX 无法打开，可能已损坏") from exc
    if "ppt/presentation.xml" not in names:
        raise AttachmentError("这不是一个 PowerPoint (.pptx) 文件")
    if total > MAX_PPTX_UNCOMPRESSED_BYTES:
        raise AttachmentError("这个 PPTX 解压后太大，无法读取")

    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise AttachmentError("这个 PPTX 无法打开，可能已损坏") from exc
    pages: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        if index > MAX_SLIDES:
            pages.append(f"（第 {MAX_SLIDES} 页之后未读取）")
            break
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_lines(shape))
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            notes = notes_frame.text.strip() if notes_frame is not None else ""
            if notes:
                lines.append(f"备注：{notes}")
        pages.append(f"第 {index} 页\n" + ("\n".join(lines) if lines else "（此页没有文字）"))
    return "\n\n".join(pages).strip()


def save_attachment(filename: str, data: bytes) -> StoredAttachment:
    """Validate and store one upload; documents are read to text right away."""

    name = _display_name(filename)
    suffix = Path(name).suffix.lower()
    if suffix in IMAGE_TYPES:
        kind, mime_type, limit = "image", IMAGE_TYPES[suffix], MAX_IMAGE_BYTES
    elif suffix in DOCUMENT_TYPES:
        kind, mime_type, limit = "document", DOCUMENT_TYPES[suffix], MAX_DOCUMENT_BYTES
    else:
        raise AttachmentError("只支持图片（PNG/JPG/WebP/GIF）、PDF 和 PPTX")
    if not data:
        raise AttachmentError("文件是空的")
    if len(data) > limit:
        raise AttachmentError(f"{'图片' if kind == 'image' else '文档'}不能超过 {limit // (1024 * 1024)} MB")
    if not _matches_signature(suffix, data):
        raise AttachmentError("文件内容和扩展名不一致")

    text = ""
    if suffix == ".pdf":
        text = _pdf_text(data)
    elif suffix == ".pptx":
        text = _pptx_text(data)

    attachment = StoredAttachment(
        id=f"{uuid.uuid4().hex}{suffix}",
        name=name,
        kind=kind,
        mime_type=mime_type,
        size=len(data),
    )
    directory = _directory()
    (directory / attachment.id).write_bytes(data)
    if kind == "document":
        (directory / f"{attachment.id}.txt").write_text(text, encoding="utf-8")
    (directory / f"{attachment.id}.json").write_text(
        json.dumps(attachment.public(), ensure_ascii=False), encoding="utf-8"
    )
    return attachment


def load_attachment(attachment_id: str) -> tuple[StoredAttachment, Path]:
    if not _ID_RE.match(attachment_id or ""):
        raise AttachmentError("附件编号无效")
    directory = _directory()
    path = directory / attachment_id
    meta_path = directory / f"{attachment_id}.json"
    if not path.is_file() or not meta_path.is_file():
        raise AttachmentError("附件不存在或已被清理，请重新添加")
    raw = json.loads(meta_path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise AttachmentError("附件信息已损坏，请重新添加")
    kind = str(raw.get("kind") or "")
    mime_type = str(raw.get("mime_type") or "")
    if kind not in {"image", "document"} or mime_type not in {*IMAGE_TYPES.values(), *DOCUMENT_TYPES.values()}:
        raise AttachmentError("附件信息已损坏，请重新添加")
    size = raw.get("size")
    attachment = StoredAttachment(
        id=attachment_id,
        name=_display_name(str(raw.get("name") or "")),
        kind=kind,
        mime_type=mime_type,
        size=size if isinstance(size, int) else path.stat().st_size,
    )
    return attachment, path


def _clip(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return text[: max(0, limit)] + _TRUNCATED


@dataclass(frozen=True)
class PreparedMessage:
    # what the model reads: the user's words plus every document's text
    model_message: str
    # what the conversation history keeps: only the words the user typed
    history_message: str
    image_ids: tuple[str, ...]
    # id / name / kind of each attachment, saved on the user message so it reloads as cards
    attachments: tuple[dict[str, str], ...] = ()


def prepare_message(message: str, attachment_ids: list[str]) -> PreparedMessage:
    if len(attachment_ids) > MAX_ATTACHMENTS:
        raise AttachmentError(f"一次最多发送 {MAX_ATTACHMENTS} 个附件")
    seen: set[str] = set()
    saved: list[dict[str, str]] = []
    images: list[str] = []
    sections: list[str] = []
    budget = MAX_TOTAL_DOCUMENT_CHARS
    for attachment_id in attachment_ids:
        if attachment_id in seen:
            continue
        seen.add(attachment_id)
        attachment, _ = load_attachment(attachment_id)
        saved.append({"id": attachment.id, "name": attachment.name, "kind": attachment.kind})
        if attachment.kind == "image":
            images.append(attachment.id)
            continue
        text_path = _directory() / f"{attachment.id}.txt"
        text = text_path.read_text(encoding="utf-8").strip() if text_path.is_file() else ""
        if not text:
            body = "（没能从这个文件读出文字，可能是纯图片或扫描件）"
        else:
            body = _clip(text, min(MAX_DOCUMENT_CHARS, budget))
            budget = max(0, budget - min(len(text), MAX_DOCUMENT_CHARS))
        sections.append(f"<attachment name=\"{attachment.name}\">\n{body}\n</attachment>")

    base = message.strip()
    model_message = base
    if sections:
        model_message = "\n\n".join(
            [base, "以下是用户随消息附上的文件内容：", *sections]
        )
    if images:
        model_message += f"\n\n（用户附上了 {len(images)} 张图片）"
    return PreparedMessage(
        model_message=model_message,
        history_message=base,
        image_ids=tuple(images),
        attachments=tuple(saved),
    )


def image_ref(attachment_id: str) -> str:
    return f"{IMAGE_REF_PREFIX}{attachment_id}"


def _image_data_url(url: str) -> str:
    attachment, path = load_attachment(url[len(IMAGE_REF_PREFIX):])
    if attachment.kind != "image":
        raise AttachmentError("附件不是图片")
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{attachment.mime_type};base64,{encoded}"


def _image_parts(llm_messages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    parts: list[dict[str, Any]] = []
    for message in llm_messages:
        content = message.get("content")
        if isinstance(content, list):
            parts.extend(
                part for part in content if isinstance(part, dict) and part.get("type") == "image_url"
            )
    return parts


def detach_images(llm_messages: list[dict[str, Any]]) -> list[str]:
    """Blank every image URL so desensitizing never touches it; returns them in order."""

    urls: list[str] = []
    for part in _image_parts(llm_messages):
        image = part.get("image_url")
        urls.append(str(image.get("url") or "") if isinstance(image, dict) else "")
        part["image_url"] = {"url": ""}
    return urls


def attach_images(llm_messages: list[dict[str, Any]], urls: list[str]) -> list[dict[str, Any]]:
    """Put the detached images back, stored references read in as data URLs."""

    parts = _image_parts(llm_messages)
    if len(parts) != len(urls):
        raise AttachmentError("图片附件在发送前丢失，请重新添加")
    for part, url in zip(parts, urls):
        if not url.startswith(IMAGE_REF_PREFIX):
            raise AttachmentError("只能发送已上传的图片附件")
        part["image_url"] = {"url": _image_data_url(url)}
    return llm_messages


def images_as_text(llm_messages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Flatten multimodal user content to text, for text-only calls such as compaction."""

    flattened: list[dict[str, Any]] = []
    for message in llm_messages:
        content = message.get("content")
        if isinstance(content, list):
            parts: list[str] = []
            for part in content:
                if isinstance(part, dict) and part.get("type") == "text":
                    parts.append(str(part.get("text") or ""))
                elif isinstance(part, dict) and part.get("type") == "image_url":
                    parts.append("[图片]")
            message = {**message, "content": "".join(parts)}
        flattened.append(message)
    return flattened


def anthropic_blocks(content: list[Any]) -> list[dict[str, Any]]:
    """Map OpenAI-style content parts to Anthropic content blocks."""

    blocks: list[dict[str, Any]] = []
    for part in content:
        if not isinstance(part, dict):
            continue
        if part.get("type") == "text" and part.get("text"):
            blocks.append({"type": "text", "text": str(part["text"])})
        elif part.get("type") == "image_url":
            image = part.get("image_url")
            url = image.get("url") if isinstance(image, dict) else ""
            match = re.match(r"^data:(image/[a-z+.-]+);base64,(.+)$", str(url or ""), re.DOTALL)
            if match is None:
                raise AttachmentError("图片附件未能转换为模型可读的格式")
            blocks.append(
                {
                    "type": "image",
                    "source": {"type": "base64", "media_type": match.group(1), "data": match.group(2)},
                }
            )
    return blocks
